# System imports
import sys
import os
import shutil
import json
import yaml

# Get the parent directory
parent_dir = os.path.abspath(os.path.join(os.path.dirname(__file__), '..'))
sys.path.insert(0, parent_dir)

# External imports
import torch
import numpy as np
import pandas as pd
from diffusers import StableDiffusionXLPipeline
import random
from PIL import Image
import matplotlib.pyplot as plt
import time
from collections import defaultdict
import csv
from pptx import Presentation
from pptx.util import Inches
from datasets import load_dataset
import clip
import argparse
import warnings

# Argument parsing
parser = argparse.ArgumentParser(description='Run random search optimization with configuration file')
parser.add_argument('--config', type=str, default="algorithms/config/config_random.yaml",
                   help='Path to configuration YAML file')
args = parser.parse_args()

config_path = args.config

# Load configuration
with open(config_path, 'r') as file:
    config = yaml.safe_load(file)

SEED = config['seed']
SEED_PATH = config['seed_path']
cuda_n = str(config['cuda'])
predictor = config['predictor']
num_inference_steps = config['num_inference_steps']
height = config['height']
width = config['width']
OUTPUT_FOLDER = config['results_folder']
NUM_ITERATIONS = config['num_iterations']
alpha = config['alpha']
beta = config['beta']
max_aesthetic_score = config['max_aesthetic_score']
max_clip_score = config['max_clip_score']
model_id = config['model_id']

# Predictor name
if predictor == 0:
    predictor_name = 'simulacra'
elif predictor == 1:
    predictor_name = 'laionv1'
elif predictor == 2:
    predictor_name = 'laionv2'
else:
    raise ValueError("Invalid predictor option.")

method_save_name = 'random_search'

# Output folder
#OUTPUT_FOLDER = f"{OUTPUT_FOLDER}/{prefix}_{method_save_name}_clip_{predictor_name}_sdxlturbo_{SEED}_a{int(alpha*100)}_b{int(beta*100)}"
OUTPUT_FOLDER = f"{OUTPUT_FOLDER}/{method_save_name}_clip_{predictor_name}_sdxlturbo_{SEED}_a{int(alpha*100)}_b{int(beta*100)}"
os.makedirs(OUTPUT_FOLDER, exist_ok=True)
shutil.copy(config_path, os.path.join(OUTPUT_FOLDER, "config_used.yaml"))

# Device
device = "cuda:" + cuda_n if torch.cuda.is_available() else "cpu"

# Pipeline
pipe = StableDiffusionXLPipeline.from_pretrained(
    model_id,
    torch_dtype=torch.float32,
    use_safetensors=True,
).to(device)
pipe.set_progress_bar_config(disable=True)

MIN_VALUE, MAX_VALUE = 0, pipe.tokenizer.vocab_size - 3
VECTOR_SIZE = pipe.tokenizer.model_max_length

# CLIP
clip_model_name = "ViT-L/14"
clip_model, clip_preprocess = clip.load(clip_model_name, device=device)

# Dataset sampling
prompt_dataset = load_dataset("nateraw/parti-prompts")["train"]
N_PER_CATEGORY = config['prompt_per_categorie']
SUBSET_SEED = config['prompt_sample_seed']
random.seed(SUBSET_SEED)

category_prompts = defaultdict(list)
for item in prompt_dataset:
    category = item.get("Category", "Uncategorized")
    category_prompts[category].append(item["Prompt"])

selected_prompts_with_category = []
for category, prompts in category_prompts.items():
    if len(prompts) >= N_PER_CATEGORY:
        sampled = random.sample(prompts, N_PER_CATEGORY)
    else:
        sampled = prompts
    for prompt in sampled:
        selected_prompts_with_category.append((prompt, category))

prompt_list_path = os.path.join(OUTPUT_FOLDER, "selected_prompts.txt")
with open(prompt_list_path, "w", encoding="utf-8") as f:
    for prompt, category in selected_prompts_with_category:
        f.write(f"{category}\t{prompt}\n")
print(f"Saved selected prompts to {prompt_list_path}")
print(f"Selected {len(selected_prompts_with_category)} prompts from {len(category_prompts)} categories.")

# Aesthetic model
if predictor == 0:
    from aesthetic_evaluation.src import simulacra_rank_image
    aesthetic_model = simulacra_rank_image.SimulacraAesthetic(device)
    model_name = 'SAM'
elif predictor == 1:
    from aesthetic_evaluation.src import laion_rank_image
    aesthetic_model = laion_rank_image.LAIONAesthetic(device, clip_model=clip_model_name)
    model_name = 'LAIONV1'
elif predictor == 2:
    from aesthetic_evaluation.src import laion_v2_rank_image
    aesthetic_model = laion_v2_rank_image.LAIONV2Aesthetic(device, clip_model=clip_model_name)
    model_name = 'LAIONV2'
else:
    raise ValueError("Invalid predictor option.")

# Seeds
if SEED_PATH is None:
    seed_list = [SEED]
else:
    with open(SEED_PATH, 'r') as file:
        seed_list = [int(line.strip()) for line in file]

def generate_image_from_prompt_tokens(token_vector, seed):
    generator = torch.Generator(device=device).manual_seed(seed)

    if isinstance(token_vector, dict):
        ids = token_vector["input_ids"]
    elif torch.is_tensor(token_vector):
        ids = token_vector
    else:
        ids = torch.tensor(token_vector, dtype=torch.long, device=device)

    ids = ids.view(1, -1).to(device)
    max_vocab = min(pipe.tokenizer.vocab_size, getattr(pipe, "tokenizer_2", pipe.tokenizer).vocab_size)
    ids = torch.clamp(ids[:, :VECTOR_SIZE], 0, max_vocab - 1)

    with torch.no_grad():
        enc_out_1 = pipe.text_encoder(ids, output_hidden_states=True)
        enc_out_2 = pipe.text_encoder_2(ids, output_hidden_states=True)

        emb_1 = enc_out_1.hidden_states[-2]
        emb_2 = enc_out_2.hidden_states[-2]
        prompt_embeds = torch.cat([emb_1, emb_2], dim=-1)
        pooled_prompt_embeds = enc_out_2[0]

    out = pipe(
        prompt_embeds=prompt_embeds,
        pooled_prompt_embeds=pooled_prompt_embeds,
        guidance_scale=0.0,
        num_inference_steps=num_inference_steps,
        generator=generator,
        output_type="pt",
    )["images"]

    image = out.clamp(0, 1).squeeze(0).permute(1, 2, 0)
    return image.to(device)

def aesthetic_evaluation(image):
    image_input = image.permute(2, 0, 1).to(torch.float32)
    score = aesthetic_model.predict_from_tensor(image_input)
    return score

def evaluate_clip_score(image_tensor, prompt):
    image = (image_tensor * 255).clamp(0, 255).byte()
    image = Image.fromarray(image.cpu().numpy())

    image_input = clip_preprocess(image).unsqueeze(0).to(device)
    text_input = clip.tokenize([prompt]).to(device)

    image_features = clip_model.encode_image(image_input)
    text_features = clip_model.encode_text(text_input)

    image_features = image_features / image_features.norm(dim=-1, keepdim=True)
    text_features = text_features / text_features.norm(dim=-1, keepdim=True)

    clip_score = (image_features @ text_features.T)
    return clip_score

def format_time(seconds):
    seconds = int(seconds)
    hours = seconds // 3600
    minutes = (seconds % 3600) // 60
    seconds = seconds % 60
    if hours > 0:
        return f"{hours}h {minutes}m {seconds}s"
    elif minutes > 0:
        return f"{minutes}m {seconds}s"
    else:
        return f"{seconds}s"

def evaluate(token_vector, seed, selected_prompt, save_path=None):
    rounded = torch.tensor(token_vector, dtype=torch.int64, device=device)
    rounded = torch.clamp(rounded, 0, pipe.tokenizer.vocab_size - 1)

    with torch.no_grad():
        image = generate_image_from_prompt_tokens(rounded, seed)
        aesthetic_score = aesthetic_evaluation(image).item()
        clip_score = evaluate_clip_score(image, selected_prompt).item()

    fitness_1 = alpha * aesthetic_score / max_aesthetic_score
    fitness_2 = beta * clip_score / max_clip_score
    fitness = fitness_1 + fitness_2

    if save_path is not None:
        image_np = image.detach().clone().cpu().numpy()
        image_np = (image_np * 255).astype(np.uint8)
        pil_image = Image.fromarray(image_np)
        pil_image.save(save_path)

    return fitness, aesthetic_score, clip_score, fitness_1, fitness_2

def detokenize(individual):
    tmp_solution = torch.tensor(individual, dtype=torch.int64)
    tmp_solution = torch.clamp(tmp_solution, 0, pipe.tokenizer.vocab_size - 1)
    decoded_string = pipe.tokenizer.decode(tmp_solution, skip_special_tokens=True, clean_up_tokenization_spaces=True)
    return decoded_string

def create_random_individual():
    return [random.randint(MIN_VALUE, MAX_VALUE) for _ in range(VECTOR_SIZE)]

# Clean any embedded NUL characters that might appear after detokenization
def _sanitize_prompt(text):
    if isinstance(text, str):
        return text.replace('\x00', '')
    return text

def main(seed, seed_number, selected_prompt, category, prompt_number):
    torch.manual_seed(seed + seed_number*1000)
    np.random.seed(seed + seed_number*1000)
    random.seed(seed + seed_number*1000)

    print(f"Selected prompt: {selected_prompt} (Category: {category})")

    results_folder = f"{OUTPUT_FOLDER}/results_{model_name}_{seed}_{prompt_number}"
    os.makedirs(results_folder, exist_ok=True)

    initial_token_vector = pipe.tokenizer(
        selected_prompt,
        padding="max_length",
        max_length=pipe.tokenizer.model_max_length,
        truncation=True,
        return_tensors="pt",
    ).to(device)

    with torch.no_grad():
        initial_image = generate_image_from_prompt_tokens(initial_token_vector["input_ids"].squeeze(0), seed)
        image_np = initial_image.detach().clone().cpu().numpy()
        image_np = (image_np * 255).astype(np.uint8)
        pil_image = Image.fromarray(image_np)
        pil_image.save(f"{results_folder}/it_0.png")

        initial_fitness, initial_aesthetic_score, initial_clip_score, initial_fitness_1, initial_fitness_2 = evaluate(initial_token_vector["input_ids"].squeeze(0), seed, selected_prompt)

    time_list = [0]

    best_fitness_overall = initial_fitness
    best_aesthetic_overall = initial_aesthetic_score
    best_clip_overall = initial_clip_score
    best_tokens_overall = initial_token_vector["input_ids"].squeeze(0).cpu().numpy()

    best_prompt_list = [selected_prompt]

    max_fit_list = [initial_fitness]
    avg_fit_list = [initial_fitness]
    std_fit_list = [0]

    max_aesthetic_score_list = [initial_aesthetic_score]
    avg_aesthetic_score_list = [initial_aesthetic_score]
    std_aesthetic_score_list = [0]

    max_clip_score_list = [initial_clip_score]
    avg_clip_score_list = [initial_clip_score]
    std_clip_score_list = [0]

    max_fitness_1_list = [initial_fitness_1]
    avg_fitness_1_list = [initial_fitness_1]
    std_fitness_1_list = [0]

    max_fitness_2_list = [initial_fitness_2]
    avg_fitness_2_list = [initial_fitness_2]
    std_fitness_2_list = [0]

    start_time = time.time()

    is_first_iteration = True
    for iteration in range(NUM_ITERATIONS):
        print(f"Iteration {iteration+1}/{NUM_ITERATIONS}")

        candidate = create_random_individual()

        fitness, aesthetic_score, clip_score, fitness_1, fitness_2 = evaluate(candidate, seed, selected_prompt)

        if fitness > best_fitness_overall or is_first_iteration:
            best_fitness_overall = fitness
            best_aesthetic_overall = aesthetic_score
            best_clip_overall = clip_score
            best_tokens_overall = np.rint(candidate).astype(np.int64)

            #with torch.no_grad():
            #    best_image = generate_image_from_prompt_tokens(torch.tensor(best_tokens_overall, dtype=torch.long, device=device), seed)
            #    image_np = best_image.detach().clone().cpu().numpy()
            #    image_np = (image_np * 255).astype(np.uint8)
            #    pil_image = Image.fromarray(image_np)
            #    pil_image.save(results_folder + "/best_%d.png" % (iteration+1))

        # Stats (single candidate per iteration)
        max_fit_list.append(best_fitness_overall)
        avg_fit_list.append(best_fitness_overall)
        std_fit_list.append(0)

        max_aesthetic_score_list.append(best_aesthetic_overall)
        avg_aesthetic_score_list.append(best_aesthetic_overall)
        std_aesthetic_score_list.append(0)

        max_clip_score_list.append(best_clip_overall)
        avg_clip_score_list.append(best_clip_overall)
        std_clip_score_list.append(0)

        #max_fitness_1_list.append(fitness_1)
        #avg_fitness_1_list.append(fitness_1)
        #std_fitness_1_list.append(0)

        #max_fitness_2_list.append(fitness_2)
        #avg_fitness_2_list.append(fitness_2)
        #std_fitness_2_list.append(0)

        best_prompt_list.append(detokenize(best_tokens_overall))

        elapsed_time = time.time() - start_time
        iterations_done = iteration + 1
        iterations_left = NUM_ITERATIONS - iterations_done
        average_time_per_iteration = elapsed_time / iterations_done
        estimated_time_remaining = average_time_per_iteration * iterations_left
        formatted_time_remaining = format_time(estimated_time_remaining)
        time_list.append(elapsed_time)

        results = pd.DataFrame({
            "iteration": list(range(0, iteration + 2)),
            "prompt": [selected_prompt] + [''] * (iteration + 1),
            "category": [category] + [''] * (iteration + 1),
            "avg_fitness": avg_fit_list,
            "std_fitness": std_fit_list,
            "max_fitness": max_fit_list,
            "avg_aesthetic_score": avg_aesthetic_score_list,
            "std_aesthetic_score": std_aesthetic_score_list,
            "max_aesthetic_score": max_aesthetic_score_list,
            "avg_clip_score": avg_clip_score_list,
            "std_clip_score": std_clip_score_list,
            "max_clip_score": max_clip_score_list,
            "best_prompt": [_sanitize_prompt(p) for p in best_prompt_list],
            "elapsed_time": time_list
        })

        results.to_csv(
            f"{results_folder}/fitness_results.csv",
            index=False,
            na_rep='nan',
            quoting=csv.QUOTE_NONNUMERIC,
            escapechar='\\'  # ensure special characters can be escaped when quoting
        )
        save_plot_results(results, results_folder)

        is_first_iteration = False

        print(f"Iteration {iteration+1}/{NUM_ITERATIONS}: Fitness: {fitness}, Aesthetic: {aesthetic_score}, CLIP: {clip_score}, Estimated time remaining: {formatted_time_remaining}")

    with torch.no_grad():
        best_image = generate_image_from_prompt_tokens(torch.tensor(best_tokens_overall, dtype=torch.long, device=device), seed)
    best_image_np = best_image.detach().cpu().numpy()
    best_image_np = (best_image_np * 255).astype(np.uint8)
    pil_image = Image.fromarray(best_image_np)
    pil_image.save(f"{results_folder}/best_all.png")

def plot_mean_std(x_axis, m_vec, std_vec, description, title=None, y_label=None, x_label=None):
    lower_bound = [M_new - Sigma for M_new, Sigma in zip(m_vec, std_vec)]
    upper_bound = [M_new + Sigma for M_new, Sigma in zip(m_vec, std_vec)]

    plt.plot(x_axis, m_vec, '--', label=description + " Avg.")
    plt.fill_between(x_axis, lower_bound, upper_bound, alpha=.3, label=description + " Avg. ± SD")
    if title is not None:
        plt.title(title)
    if y_label is not None:
        plt.ylabel(y_label)
    if x_label is not None:
        plt.xlabel(x_label)

def save_plot_results(results, results_folder):
    plt.figure(figsize=(10, 6))
    plot_mean_std(results['iteration'], results['avg_fitness'], results['std_fitness'], "Fitness")
    plt.plot(results['iteration'], results['max_fitness'], 'r-', label="Fitness")
    plt.ylim(0, 1.1)
    plt.xlabel('Iteration')
    plt.ylabel('Fitness')
    plt.grid()
    plt.legend(loc="upper left", bbox_to_anchor=(1, 1))
    plt.tight_layout()
    plt.savefig(results_folder + "/fitness_evolution.png")
    plt.close()

    plt.figure(figsize=(10, 6))
    plot_mean_std(results['iteration'], results['avg_aesthetic_score'], results['std_aesthetic_score'], "Population")
    plt.plot(results['iteration'], results['max_aesthetic_score'], 'r-', label="Best")
    plt.ylim(0, 10)
    plt.xlabel('Iteration')
    plt.ylabel('Aesthetic Score')
    plt.grid()
    plt.legend(loc="upper left", bbox_to_anchor=(1, 1))
    plt.tight_layout()
    plt.savefig(results_folder + "/aesthetic_score_evolution.png")
    plt.close()

    plt.figure(figsize=(10, 6))
    plot_mean_std(results['iteration'], results['avg_clip_score'], results['std_clip_score'], "Population")
    plt.plot(results['iteration'], results['max_clip_score'], 'r-', label="Best")
    plt.ylim(0, 0.6)
    plt.xlabel('Iteration')
    plt.ylabel('CLIP Score')
    plt.grid()
    plt.legend(loc="upper left", bbox_to_anchor=(1, 1))
    plt.tight_layout()
    plt.savefig(results_folder + "/clip_score_evolution.png")
    plt.close()

def aggregate_results():
    aggregated_data = None

    for folder_name in os.listdir(OUTPUT_FOLDER):
        if folder_name.startswith(f"results_{model_name}_"):
            seed = folder_name.split("_")[-2]
            prompt_number = folder_name.split("_")[-1]

            file_path = os.path.join(OUTPUT_FOLDER, folder_name, "fitness_results.csv")

            if os.path.exists(file_path):
                df = pd.read_csv(file_path)
                df = df.drop(columns=["prompt", "category", "best_prompt"])

                df = df.rename(columns={"avg_fitness": f"avg_fitness_{seed}_{prompt_number}"})
                df = df.rename(columns={"max_fitness": f"max_fitness_{seed}_{prompt_number}"})
                df = df.rename(columns={"std_fitness": f"std_fitness_{seed}_{prompt_number}"})
                df = df.rename(columns={"avg_aesthetic_score": f"avg_aesthetic_score_{seed}_{prompt_number}"})
                df = df.rename(columns={"max_aesthetic_score": f"max_aesthetic_score_{seed}_{prompt_number}"})
                df = df.rename(columns={"std_aesthetic_score": f"std_aesthetic_score_{seed}_{prompt_number}"})
                df = df.rename(columns={"avg_clip_score": f"avg_clip_score_{seed}_{prompt_number}"})
                df = df.rename(columns={"max_clip_score": f"max_clip_score_{seed}_{prompt_number}"})
                df = df.rename(columns={"std_clip_score": f"std_clip_score_{seed}_{prompt_number}"})
                df = df.rename(columns={"elapsed_time": f"elapsed_time_{seed}_{prompt_number}"})

                if aggregated_data is None:
                    aggregated_data = df
                else:
                    aggregated_data = pd.merge(aggregated_data, df, on="iteration", how="outer")
            else:
                print(f"File not found: {file_path}")

    if aggregated_data is None:
        print("No data was aggregated. Check the input folders and files.")
        return

    output_file = os.path.join(OUTPUT_FOLDER, "aggregated_score_results.xlsx")
    aggregated_data.to_excel(output_file, index=False)
    print(f"Aggregated results saved to {output_file}")

    data = pd.read_excel(output_file)

    data['avg_fitness'] = data.filter(like='avg_fitness_').mean(axis=1)
    data['std_fitness'] = data.filter(like='std_fitness_').std(axis=1)
    data['best_avg_fitness'] = data.filter(like='max_fitness_').mean(axis=1)
    data['best_std_fitness'] = data.filter(like='max_fitness_').std(axis=1)
    data['max_fitness'] = data.filter(like='max_fitness_').max(axis=1)
    data['avg_aesthetic_score'] = data.filter(like='avg_aesthetic_score_').mean(axis=1)
    data['std_aesthetic_score'] = data.filter(like='std_aesthetic_score_').std(axis=1)
    data['best_avg_aesthetic_score'] = data.filter(like='max_aesthetic_score_').mean(axis=1)
    data['best_std_aesthetic_score'] = data.filter(like='max_aesthetic_score_').std(axis=1)
    data['max_aesthetic_score'] = data.filter(like='max_aesthetic_score_').max(axis=1)
    data['avg_clip_score'] = data.filter(like='avg_clip_score_').mean(axis=1)
    data['std_clip_score'] = data.filter(like='std_clip_score_').std(axis=1)
    data['best_avg_clip_score'] = data.filter(like='max_clip_score_').mean(axis=1)
    data['best_std_clip_score'] = data.filter(like='max_clip_score_').std(axis=1)
    data['max_clip_score'] = data.filter(like='max_clip_score_').max(axis=1)

    plt.figure(figsize=(10, 6))
    plot_mean_std(data['iteration'], data['avg_fitness'], data['std_fitness'], "Population")
    plot_mean_std(data['iteration'], data['best_avg_fitness'], data['best_std_fitness'], "Bests")
    plt.plot(data['iteration'], data['max_fitness'], 'r-', label="Best")
    plt.ylim(0, 1.1)
    plt.xlabel('Iteration')
    plt.ylabel('Fitness')
    plt.grid()
    plt.legend(loc="upper left", bbox_to_anchor=(1, 1))
    plt.tight_layout()
    plt.savefig(OUTPUT_FOLDER + "/fitness_evolution.png")
    plt.close()

    plt.figure(figsize=(10, 6))
    plot_mean_std(data['iteration'], data['avg_aesthetic_score'], data['std_aesthetic_score'], "Population")
    plot_mean_std(data['iteration'], data['best_avg_aesthetic_score'], data['best_std_aesthetic_score'], "Bests")
    plt.plot(data['iteration'], data['max_aesthetic_score'], 'r-', label="Best")
    plt.ylim(0, 10.5)
    plt.xlabel('Iteration')
    plt.ylabel('Aesthetic Score')
    plt.grid()
    plt.legend(loc="upper left", bbox_to_anchor=(1, 1))
    plt.tight_layout()
    plt.savefig(OUTPUT_FOLDER + "/aesthetic_score_evolution.png")
    plt.close()

    plt.figure(figsize=(10, 6))
    plot_mean_std(data['iteration'], data['avg_clip_score'], data['std_clip_score'], "Population")
    plot_mean_std(data['iteration'], data['best_avg_clip_score'], data['best_std_clip_score'], "Bests")
    plt.plot(data['iteration'], data['max_clip_score'], 'r-', label="Best")
    plt.xlabel('Iteration')
    plt.ylabel('CLIP Score')
    plt.ylim(0, 0.6)
    plt.grid()
    plt.legend(loc="upper left", bbox_to_anchor=(1, 1))
    plt.tight_layout()
    plt.savefig(OUTPUT_FOLDER + "/clip_score_evolution.png")
    plt.close()

    presentation = Presentation()

    folders = []
    for folder_name in os.listdir(OUTPUT_FOLDER):
        if folder_name.startswith("results_"):
            seed_number = int(folder_name.split("_")[-2])
            prompt_number = int(folder_name.split("_")[-1])
            folder_path = os.path.join(OUTPUT_FOLDER, folder_name)
            folders.append((seed_number, prompt_number, folder_path))

    folders.sort(key=lambda x: x[1])

    for seed_number, prompt_number, folder_path in folders:
        it_0_path = os.path.join(folder_path, "it_0.png")
        best_all_path = os.path.join(folder_path, "best_all.png")
        fitness_evolution_path = os.path.join(folder_path, "fitness_evolution.png")
        aesthetic_evolution_path = os.path.join(folder_path, "aesthetic_score_evolution.png")
        clip_evolution_path = os.path.join(folder_path, "clip_score_evolution.png")
        csv_path = os.path.join(folder_path, "fitness_results.csv")

        fitness_initial = None
        fitness_best = None
        aesthetic_initial = None
        aesthetic_best = None
        clip_initial = None
        clip_best = None
        prompt_text = None
        category = None

        if os.path.exists(csv_path):
            with open(csv_path, 'r') as csvfile:
                reader = csv.DictReader(csvfile)
                rows = list(reader)
                if rows:
                    first_row = rows[0]
                    fitness_initial = float(first_row['max_fitness'])
                    aesthetic_initial = float(first_row['max_aesthetic_score'])
                    clip_initial = float(first_row['max_clip_score'])
                    prompt_text = first_row['prompt']
                    category = first_row['category']

                    best_row = max(rows, key=lambda r: float(r['max_fitness']))
                    fitness_best = float(best_row['max_fitness'])
                    aesthetic_best = float(best_row['max_aesthetic_score'])
                    clip_best = float(best_row['max_clip_score'])

        if os.path.exists(it_0_path) and os.path.exists(best_all_path):
            slide = presentation.slides.add_slide(presentation.slide_layouts[5])
            title = slide.shapes.title
            title.text = f"Seed {seed_number}"

            if prompt_text:
                left = Inches(0.5)
                top = Inches(1)
                width = Inches(9)
                textbox = slide.shapes.add_textbox(left, top, width, Inches(0.5))
                textbox.text = f"Prompt: {prompt_text}"

            if category:
                left = Inches(0.5)
                top = Inches(1.5)
                width = Inches(9)
                textbox = slide.shapes.add_textbox(left, top, width, Inches(0.5))
                textbox.text = f"Category: {category}"

            slide.shapes.add_picture(it_0_path, Inches(0.5), Inches(2), height=Inches(4))

            left = Inches(0.5)
            top = Inches(6.2)
            width = Inches(4)
            textbox = slide.shapes.add_textbox(left, top, width, Inches(0.5))
            text = "Initial iteration"
            if fitness_initial is not None:
                text += f"\nInitial Fitness: {fitness_initial:.4f}"
            if aesthetic_initial is not None:
                text += f"\nAesthetic Score: {aesthetic_initial:.4f}"
            if clip_initial is not None:
                text += f"\nCLIP Score: {clip_initial:.4f}"
            textbox.text = text

            slide.shapes.add_picture(best_all_path, Inches(5.5), Inches(2), height=Inches(4))

            left = Inches(5.5)
            top = Inches(6.2)
            width = Inches(4)
            textbox = slide.shapes.add_textbox(left, top, width, Inches(0.5))
            text = "Best iteration"
            if fitness_best is not None:
                text += f"\nBest fitness: {fitness_best:.4f}"
            if aesthetic_best is not None:
                text += f"\nAesthetic Score: {aesthetic_best:.4f}"
            if clip_best is not None:
                text += f"\nCLIP Score: {clip_best:.4f}"
            textbox.text = text

        if os.path.exists(fitness_evolution_path):
            slide = presentation.slides.add_slide(presentation.slide_layouts[5])
            title = slide.shapes.title
            title.text = f"Seed {seed_number}"
            slide.shapes.add_picture(fitness_evolution_path, Inches(0), Inches(2), height=Inches(4))

            left = Inches(0.5)
            top = Inches(6.2)
            width = Inches(4)
            textbox = slide.shapes.add_textbox(left, top, width, Inches(0.5))
            textbox.text = "Fitness evolution"

        if os.path.exists(clip_evolution_path) and os.path.exists(aesthetic_evolution_path):
            slide = presentation.slides.add_slide(presentation.slide_layouts[5])
            title = slide.shapes.title
            title.text = f"Seed {seed_number}"

            slide.shapes.add_picture(clip_evolution_path, Inches(0), Inches(2), height=Inches(4))

            left = Inches(0.5)
            top = Inches(6.2)
            width = Inches(4)
            textbox = slide.shapes.add_textbox(left, top, width, Inches(0.5))
            textbox.text = "CLIP score evolution"

            slide.shapes.add_picture(aesthetic_evolution_path, Inches(5), Inches(2), height=Inches(4))

            left = Inches(5.5)
            top = Inches(6.2)
            width = Inches(4)
            textbox = slide.shapes.add_textbox(left, top, width, Inches(0.5))
            textbox.text = "Aesthetic score evolution"

    output_filename = os.path.join(OUTPUT_FOLDER, f"summary.pptx")
    presentation.save(output_filename)
    print(f"Presentation saved as {output_filename}")

if __name__ == "__main__":
    seed_number = 1
    for seed in seed_list:
        prompt_number = 1
        for prompt, category in selected_prompts_with_category:
            print(f"Running seed {seed}, prompt: {prompt} (Category: {category})")
            main(seed, seed_number, prompt, category, prompt_number)
            print(f"Run with seed {seed} and prompt '{prompt}' finished!")
            aggregate_results()
            prompt_number += 1
        seed_number += 1
        #aggregate_results()
